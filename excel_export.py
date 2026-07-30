"""云端可部署的教师教学工作量 Excel 导出模块。"""

from __future__ import annotations

import io
import math
import re
import unicodedata
import zipfile
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins


FONT_NAME = "SimHei"
EDIT_PASSWORD = "20130052"
DATA_COLUMN_WIDTHS = [16, 58, 10, 13, 32, 26]
MIN_DATA_ROW_HEIGHT = 36
THIN_BLACK = Side(style="thin", color="000000")
ALL_BORDERS = Border(
    left=THIN_BLACK,
    right=THIN_BLACK,
    top=THIN_BLACK,
    bottom=THIN_BLACK,
)


def _display_width(text: object) -> int:
    """按中英文字符的近似显示宽度计算 Excel 换行占用。"""
    return sum(
        2 if unicodedata.east_asian_width(character) in {"W", "F", "A"} else 1
        for character in str(text or "")
    )


def _wrapped_line_count(text: object, column_width: float) -> int:
    """估算启用自动换行后文本在指定列宽中需要的行数。"""
    usable_width = max(int(column_width) - 1, 1)
    paragraphs = str(text or "").splitlines() or [""]
    return sum(
        max(1, math.ceil(_display_width(paragraph) / usable_width))
        for paragraph in paragraphs
    )


def _content_row_height(values: list[object]) -> float:
    """以现有 36 磅为最小值，按最长单元格内容自动增加行高。"""
    required_lines = max(
        _wrapped_line_count(value, DATA_COLUMN_WIDTHS[index])
        for index, value in enumerate(values)
    )
    estimated_height = required_lines * 13.5 + 8
    return min(409, max(MIN_DATA_ROW_HEIGHT, estimated_height))


def _patch_basic_info_rich_text(xlsx_bytes: bytes, data: dict) -> bytes:
    """为基本信息中的用户输入值增加下划线。"""
    source = io.BytesIO(xlsx_bytes)
    target = io.BytesIO()
    runs = [
        ("所属系部：", data.get("dept", "")),
        ("    姓名：", data.get("name", "")),
        ("    职称：", data.get("title", "")),
        ("    填报日期：", data.get("fill_date", "")),
    ]
    rich_text = "".join(
        (
            '<r><rPr><rFont val="SimHei"/><sz val="11"/>'
            '<color rgb="FF000000"/></rPr>'
            f'<t xml:space="preserve">{escape(str(label))}</t></r>'
            '<r><rPr><rFont val="SimHei"/><sz val="11"/>'
            '<color rgb="FF000000"/><u/></rPr>'
            f'<t xml:space="preserve">{escape(str(value or ""))}</t></r>'
        )
        for label, value in runs
    )

    with zipfile.ZipFile(source, "r") as input_zip:
        with zipfile.ZipFile(
            target, "w", compression=zipfile.ZIP_DEFLATED
        ) as output_zip:
            for item in input_zip.infolist():
                payload = input_zip.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    sheet_xml = payload.decode("utf-8")

                    def replace_cell(match: re.Match) -> str:
                        attributes = match.group(1)
                        style_match = re.search(r'\bs="[^"]+"', attributes)
                        style_attribute = (
                            f" {style_match.group(0)}" if style_match else ""
                        )
                        return (
                            f'<c r="A3"{style_attribute} t="inlineStr">'
                            f"<is>{rich_text}</is></c>"
                        )

                    sheet_xml = re.sub(
                        r'<c r="A3"([^>]*)>.*?</c>',
                        replace_cell,
                        sheet_xml,
                        count=1,
                    )
                    payload = sheet_xml.encode("utf-8")
                output_zip.writestr(item, payload)
    return target.getvalue()


def build_workload_excel(data: dict) -> bytes:
    """根据已排序的导出数据生成 Excel 文件字节。"""
    rows = data.get("rows") if isinstance(data.get("rows"), list) else []
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "工作量确认表"
    sheet.sheet_view.showGridLines = False

    term = rows[0].get("term", "") if rows else ""
    sheet.merge_cells("A1:F1")
    sheet["A1"] = f"科技商学院{term}教师教学工作量填报表"
    sheet["A1"].font = Font(name=FONT_NAME, size=16, bold=False)
    sheet["A1"].alignment = Alignment(
        horizontal="center", vertical="center"
    )
    sheet.row_dimensions[1].height = 28

    sheet.merge_cells("A3:F3")
    sheet["A3"] = (
        f"所属系部：{data.get('dept', '')}    "
        f"姓名：{data.get('name', '')}    "
        f"职称：{data.get('title', '')}    "
        f"填报日期：{data.get('fill_date', '')}"
    )
    sheet["A3"].font = Font(name=FONT_NAME, size=11, bold=False)
    sheet["A3"].alignment = Alignment(
        horizontal="center", vertical="center"
    )

    headers = ["项目类别", "计算标准", "工作量", "完成时间", "完成说明", "审核人"]
    for column, value in enumerate(headers, start=1):
        cell = sheet.cell(row=5, column=column, value=value)
        cell.font = Font(name=FONT_NAME, size=10, bold=False)
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        cell.border = ALL_BORDERS
    sheet.row_dimensions[5].height = 22

    first_data_row = 6
    export_rows = rows or [
        {
            "category": "",
            "standard": "",
            "workload": 0,
            "time": "",
            "remark": "",
            "reviewer": "",
        }
    ]
    for row_index, row in enumerate(export_rows, start=first_data_row):
        values = [
            row.get("category", ""),
            row.get("standard", ""),
            float(row.get("workload") or 0),
            row.get("time", ""),
            row.get("remark", ""),
            row.get("reviewer", ""),
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=column, value=value)
            cell.font = Font(name=FONT_NAME, size=9, bold=False)
            cell.alignment = Alignment(
                horizontal="center" if column in (3, 4) else "left",
                vertical="center",
                wrap_text=True,
            )
            cell.border = ALL_BORDERS
        sheet.cell(row=row_index, column=3).number_format = "General"
        sheet.cell(row=row_index, column=4).number_format = "yyyy-mm-dd"
        sheet.row_dimensions[row_index].height = _content_row_height(values)

    last_data_row = first_data_row + len(export_rows) - 1
    total_row = last_data_row + 2
    sheet.cell(row=total_row, column=4, value="工作量合计：")
    sheet.merge_cells(
        start_row=total_row,
        start_column=5,
        end_row=total_row,
        end_column=6,
    )
    sheet.cell(
        row=total_row,
        column=5,
        value=f'=SUM(C{first_data_row}:C{last_data_row})&" 学时"',
    )
    for column in range(4, 7):
        cell = sheet.cell(row=total_row, column=column)
        cell.font = Font(name=FONT_NAME, size=11, bold=False)
        cell.alignment = Alignment(horizontal="right", vertical="center")

    sign_row = total_row + 2
    sheet.merge_cells(
        start_row=sign_row, start_column=1, end_row=sign_row, end_column=2
    )
    sheet.merge_cells(
        start_row=sign_row, start_column=3, end_row=sign_row, end_column=6
    )
    sheet.cell(
        row=sign_row, column=1, value="教师本人签字：________________"
    )
    sheet.cell(
        row=sign_row, column=3, value="系主任审核签字：________________"
    )
    for column in range(1, 7):
        cell = sheet.cell(row=sign_row, column=column)
        cell.font = Font(name=FONT_NAME, size=11, bold=False)
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[sign_row].height = 30

    for column, width in enumerate(DATA_COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width

    sheet.freeze_panes = "A6"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.orientation = sheet.ORIENTATION_LANDSCAPE
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True
    sheet.page_margins = PageMargins(
        left=0.25,
        right=0.25,
        top=0.35,
        bottom=0.35,
        header=0.1,
        footer=0.1,
    )
    sheet.print_area = f"A1:F{sign_row}"
    sheet.print_title_rows = "5:5"
    sheet.protection.set_password(EDIT_PASSWORD)
    sheet.protection.sheet = True
    try:
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
    except AttributeError:
        pass

    raw_output = io.BytesIO()
    workbook.save(raw_output)
    return _patch_basic_info_rich_text(raw_output.getvalue(), data)
